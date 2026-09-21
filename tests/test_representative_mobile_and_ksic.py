"""Synthetic fixtures only; contact data never reaches a network service."""
from io import BytesIO
import ast
import json
from pathlib import Path
from unittest.mock import Mock

import pandas as pd
import pytest

import cloud_restore
import cloud_sync
import enterprise_customer_management as directory
import enterprise_registration_ui as registration
import registered_policy_match
import utils


BUSINESS = "-".join(("111", "11", "11111"))
MOBILE = "-".join(("010", "0000", "0000"))
OTHER_MOBILE = "-".join(("011", "000", "0000"))
FIELDS = ("대표자 휴대전화", "표준산업분류코드", "표준산업분류차수")


def company(**overrides):
    return {"업체명": "합성 검증기업", "대표자명": "검증대표", "사업자등록번호": BUSINESS,
            "대표자 휴대전화": MOBILE, "표준산업분류코드": "C24321",
            "표준산업분류차수": "11차", "주업종코드": "123456", **overrides}


@pytest.fixture
def storage(tmp_path, monkeypatch):
    path = tmp_path / "synthetic-customers.xlsx"
    monkeypatch.setattr(utils, "get_user_cumulative_db_path", lambda _: path)
    monkeypatch.setattr(utils, "find_customer_template", lambda: None)
    return path


@pytest.mark.parametrize("value,expected", [
    (MOBILE, MOBILE), (MOBILE.replace("-", ""), MOBILE),
    ("+82 " + MOBILE[1:].replace("-", " "), MOBILE),
    (OTHER_MOBILE, OTHER_MOBILE),
    ("-".join(("019", "0000", "0000")), "-".join(("019", "0000", "0000"))),
    ("", ""), (None, ""), (pd.NA, ""),
])
def test_manual_mobile_normalization(value, expected):
    assert utils.normalize_representative_mobile(value) == expected


@pytest.mark.parametrize("value", [
    "-".join(("02", "0000", "0000")), "123", "문자" + MOBILE,
    "-".join(("010", "000", "0000")), MOBILE + "/" + OTHER_MOBILE,
])
def test_invalid_mobile_cannot_save(value, storage):
    with pytest.raises(ValueError, match="대표자 휴대전화"):
        utils.normalize_representative_mobile(value)
    _, error = registration.prepare_registration_data(company(), {"대표자 휴대전화": value})
    assert "대표자 휴대전화" in error
    utils.append_cretop_to_user_customer_db("unused.pdf", "synthetic-owner", extracted_data=company())
    before = storage.read_bytes()
    ok, message = utils.update_user_customer_record("synthetic-owner", 0, {"대표자 휴대전화": value})
    assert not ok and "대표자 휴대전화" in message
    assert storage.read_bytes() == before


def test_review_does_not_infer_representative_mobile_from_other_contact_fields():
    data = {key: value for key, value in company().items() if key != "대표자 휴대전화"}
    data.update({"휴대전화": MOBILE, "대표전화": MOBILE, "간편인증휴대전화": MOBILE})
    prepared, error = registration.prepare_registration_data(data, {})
    assert not error
    assert prepared["대표자 휴대전화"] == ""
    prepared, error = registration.prepare_registration_data(company(), {"대표자 휴대전화": ""})
    assert not error and prepared["대표자 휴대전화"] == MOBILE


def test_new_fields_survive_legacy_template_and_download_template(tmp_path, monkeypatch):
    template = tmp_path / "synthetic-template.xlsx"
    pd.DataFrame(columns=["업체명", "대표자명", "기존확장열"]).to_excel(template, sheet_name="고객DB", index=False)
    monkeypatch.setattr(utils, "find_customer_template", lambda: template)
    columns = utils.get_customer_db_columns()
    assert set(FIELDS) <= set(columns)
    assert "기존확장열" in columns
    template_bytes = utils.make_basic_customer_template_bytes()
    headers = pd.read_excel(BytesIO(template_bytes), sheet_name="고객DB", nrows=0).columns
    assert set(FIELDS) <= set(headers)


def test_customer_excel_round_trip_and_reviewed_update_preserve_tax_code(storage):
    data = company(**{"대표자 휴대전화": MOBILE.replace("-", ""), "표준산업분류코드": "01111"})
    _, count, _, _, _ = utils.append_cretop_to_user_customer_db("unused.pdf", "synthetic-owner", extracted_data=data)
    assert count == 1
    first = utils._read_cumulative_customer_db(storage).iloc[0]
    assert first["대표자 휴대전화"] == data["대표자 휴대전화"]
    assert first["표준산업분류코드"] == "01111"
    reviewed = company(**{"대표자 휴대전화": OTHER_MOBILE, "표준산업분류코드": "C24322"})
    reviewed.pop("주업종코드")
    ok, _, _ = utils.refresh_existing_customer_from_cretop("synthetic-owner", reviewed, reviewed_fields=registration.BASIC_FIELDS)
    assert ok
    saved = utils._read_cumulative_customer_db(storage).iloc[0]
    assert saved["대표자 휴대전화"] == OTHER_MOBILE
    assert saved["표준산업분류코드"] == "C24322"
    assert str(saved["주업종코드"]) == "123456"
    blank = {**reviewed, **{key: "" for key in FIELDS}}
    assert utils.refresh_existing_customer_from_cretop("synthetic-owner", blank, reviewed_fields=registration.BASIC_FIELDS)[0]
    saved_again = utils._read_cumulative_customer_db(storage).iloc[0]
    assert all(saved_again[key] == saved[key] for key in FIELDS)


def test_direct_customer_edit_normalizes_mobile_and_blank_preserves_existing(storage):
    utils.append_cretop_to_user_customer_db("unused.pdf", "synthetic-owner", extracted_data=company())
    assert utils.update_user_customer_record("synthetic-owner", 0, {"대표자 휴대전화": OTHER_MOBILE.replace("-", "")})[0]
    assert utils._read_cumulative_customer_db(storage).iloc[0]["대표자 휴대전화"] == OTHER_MOBILE
    assert utils.update_user_customer_record("synthetic-owner", 0, {"대표자 휴대전화": "", "대표자명": "수정대표"})[0]
    assert utils._read_cumulative_customer_db(storage).iloc[0]["대표자 휴대전화"] == OTHER_MOBILE


def test_excel_upload_keeps_leading_zero_fields(storage, tmp_path):
    source = tmp_path / "synthetic-upload.xlsx"
    data = company(**{"대표자 휴대전화": MOBILE.replace("-", ""), "표준산업분류코드": "01111"})
    pd.DataFrame([data]).to_excel(source, sheet_name="고객DB", index=False)
    _, count = utils.append_user_customer_db(source, "synthetic-owner")
    assert count == 1
    saved = utils._read_cumulative_customer_db(storage).iloc[0]
    assert saved["대표자 휴대전화"] == data["대표자 휴대전화"]
    assert saved["표준산업분류코드"] == "01111"
    loaded = registered_policy_match._load_registered_customers_from_excel(storage).iloc[0]
    assert loaded["대표자 휴대전화"] == data["대표자 휴대전화"]
    assert loaded["표준산업분류코드"] == "01111"


def test_customer_json_payload_and_cloud_restore_preserve_independent_fields(storage, monkeypatch):
    data = company()
    parameters = cloud_sync._customer_profile_parameters(
        "synthetic-owner", data, business_no=BUSINESS, source="synthetic", manager_name="", customer_id="", previous_business_no="",
    )
    transmitted = json.loads(json.dumps(parameters, ensure_ascii=False))["p_customer_data"]
    assert all(transmitted[key] == data[key] for key in FIELDS)
    db = Mock()
    db.select_all.return_value = [{"customer_data": transmitted, "business_no": BUSINESS}]
    monkeypatch.setattr(cloud_restore, "CloudDatabase", lambda: db)
    monkeypatch.setattr(cloud_restore, "cloud_is_configured", lambda: True)
    monkeypatch.setattr(cloud_restore, "get_user_cumulative_db_path", lambda _: storage)
    monkeypatch.setattr(cloud_restore, "find_customer_template", lambda: None)
    result = cloud_restore.restore_customer_db_if_needed("synthetic-owner")
    assert result["restored"]
    restored = utils._read_cumulative_customer_db(storage).iloc[0]
    assert all(restored[key] == data[key] for key in FIELDS)


def test_all_customer_search_paths_include_representative_mobile(monkeypatch):
    frame = pd.DataFrame([company(), {"업체명": "다른 합성기업"}]).fillna("")
    monkeypatch.setattr(directory, "active_customers", lambda *_: (frame, {}))
    monkeypatch.setattr(directory, "get_crm_file_path", lambda _: None)
    monkeypatch.setattr(directory, "_profile_map", lambda _: {})
    for query in (MOBILE, MOBILE.replace("-", "")):
        assert len(directory.search_customer_rows(frame, query)) == 1
        assert len(directory.filter_active_customers("synthetic-owner", frame, query)) == 1
    rows = directory._customer_directory("synthetic-owner", frame)
    assert MOBILE in rows[0]["search_text"]
    assert MOBILE.replace("-", "") in rows[0]["search_text"]


def test_basic_details_and_manual_edit_expose_independent_fields():
    root = Path(__file__).resolve().parents[1]
    app_source = (root / "app.py").read_text(encoding="utf-8")
    center_source = (root / "enterprise_center.py").read_text(encoding="utf-8")
    for key in FIELDS:
        assert f'"{key}"' in center_source
        assert f'"{key}"' in app_source
    assert 'key=f"edit_representative_mobile_{selected_idx}"' in app_source
    assert 'update_values["대표자 휴대전화"] = representative_mobile' in app_source


def test_customer_management_search_keeps_all_columns_and_accepts_phone_digits():
    # Compile the pure function without importing app.py's authentication/UI side effects.
    root = Path(__file__).resolve().parents[1]
    tree = ast.parse((root / "app.py").read_text(encoding="utf-8"))
    function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "_filter_customer_management_rows")
    namespace = {}
    exec(compile(ast.Module(body=[function], type_ignores=[]), "app.py", "exec"), namespace)
    search = namespace[function.name]
    frame = pd.DataFrame([company(), {"업체명": "다른 합성기업", "사용자확장메모": "확장검색"}]).fillna("")
    assert len(search(frame, "")) == 2
    assert len(search(frame, MOBILE)) == 1
    assert len(search(frame, MOBILE.replace("-", ""))) == 1
    assert search(frame, "확장검색").iloc[0]["업체명"] == "다른 합성기업"


def test_legacy_full_extraction_reuses_identity_and_industry_without_changing_financials(monkeypatch):
    text = "\n".join((
        "기업명", "사업자번호", "보고서 표지", "기업명 합성 검증기업", "영문기업명 SYNTHETIC",
        "사업자번호 " + BUSINESS, "대표자명 검증대표", "종업원수 3명",
        "표준산업분류(10차) (C24320) 합성 이전업종", "표준산업분류(11차) (C24321) 합성 최신업종",
        "주요제품 검증제품",
    ))
    monkeypatch.setattr(utils, "extract_pdf_text", lambda _: (text, ""))
    amount = Mock(return_value=123)
    monkeypatch.setattr(utils, "_latest_summary_amount_million", amount)
    data, error = utils.extract_cretop_pdf_data("unused.pdf")
    assert not error
    assert data["업체명"] == "합성 검증기업"
    assert data["대표자명"] == "검증대표"
    assert data["표준산업분류코드"] == "C24321"
    assert data["표준산업분류차수"] == "11차"
    assert data["업종명"] == "합성 최신업종"
    assert "주업종코드" not in data
    assert data["매출액"] == data["영업이익"] == 123
    assert amount.call_count == 6


@pytest.mark.parametrize("name", ["사업자번호", "기업명", "-", "기업명 사업자번호"])
def test_label_names_cannot_bypass_review_or_direct_append(name, storage):
    data = company(업체명=name)
    assert registration.prepare_registration_data(data, {})[1]
    _, count, _, _, _ = utils.append_cretop_to_user_customer_db("unused.pdf", "synthetic-owner", extracted_data=data)
    assert count == 0
    assert not storage.exists()


@pytest.fixture
def legacy_workbook_with_extensions(storage):
    """A valid old workbook made from scratch; never use a customer/template file."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    old_columns = [column for column in utils.get_customer_db_columns() if column not in FIELDS]
    columns = old_columns + ["사용자확장메모", "사용자계산식"]
    row = {**company(), "매출액": "=10+20", "사용자확장메모": "합성 보존 메모", "사용자계산식": "=1+2"}
    workbook = Workbook()
    customer_sheet = workbook.active
    customer_sheet.title = "고객DB"
    customer_sheet.append(columns)
    customer_sheet.append([row.get(column, "") for column in columns])
    customer_sheet.cell(2, columns.index("사용자확장메모") + 1).comment = Comment("합성 보존 주석", "합성 작성자")
    for name in ("상시정책자금DB", "고용지원금DB", "코드표", "사용가이드", "사용자추가시트"):
        worksheet = workbook.create_sheet(name)
        worksheet["A1"] = "합성 보존값 " + name
        worksheet["B2"] = "=1+2"
    workbook.save(storage)
    workbook.close()

    # It was current immediately before the three new columns were required.
    stat = storage.stat()
    assert utils._inspect_cumulative_customer_db_format(
        str(storage), stat.st_mtime_ns, stat.st_size, tuple(old_columns),
    ) == (True, 1)
    return storage


def test_new_contact_columns_migration_preserves_existing_customer_extensions(legacy_workbook_with_extensions):
    from openpyxl import load_workbook

    before = legacy_workbook_with_extensions.read_bytes()
    previous_mtime = legacy_workbook_with_extensions.stat().st_mtime_ns
    path, count, converted = utils.ensure_user_cumulative_db_format("synthetic-owner")
    assert path == legacy_workbook_with_extensions and count == 1
    if not converted:
        assert path.read_bytes() == before
        assert path.stat().st_mtime_ns == previous_mtime
    with path.open("rb") as source:
        workbook = load_workbook(source, data_only=False)
        try:
            worksheet = workbook["고객DB"]
            headers = [cell.value for cell in worksheet[1]]
            if converted:
                assert set(FIELDS) <= set(headers)
            assert "사용자확장메모" in headers
            assert "사용자계산식" in headers
            assert worksheet.cell(2, headers.index("사용자확장메모") + 1).value == "합성 보존 메모"
            assert worksheet.cell(2, headers.index("사용자계산식") + 1).value == "=1+2"
            assert str(worksheet.cell(2, headers.index("주업종코드") + 1).value) == "123456"
            assert worksheet.cell(2, headers.index("매출액") + 1).value == "=10+20"
            assert worksheet.cell(2, headers.index("사용자확장메모") + 1).comment.text == "합성 보존 주석"
        finally:
            workbook.close()


@pytest.mark.parametrize("sheet_name", ["상시정책자금DB", "사용가이드", "사용자추가시트"])
def test_new_contact_columns_migration_preserves_existing_other_sheets(legacy_workbook_with_extensions, sheet_name):
    from openpyxl import load_workbook

    path, count, converted = utils.ensure_user_cumulative_db_format("synthetic-owner")
    assert path == legacy_workbook_with_extensions and count == 1
    with path.open("rb") as source:
        workbook = load_workbook(source, data_only=False)
        try:
            assert sheet_name in workbook.sheetnames
            assert workbook[sheet_name]["A1"].value == "합성 보존값 " + sheet_name
            assert workbook[sheet_name]["B2"].value == "=1+2"
        finally:
            workbook.close()


def _assert_existing_workbook_data_preserved(path, *, expected_business_no=BUSINESS, expected_rows=1):
    from openpyxl import load_workbook

    original_headers = [column for column in utils.get_customer_db_columns() if column not in FIELDS]
    original_headers += ["사용자확장메모", "사용자계산식"]
    workbook = load_workbook(path, data_only=False)
    try:
        worksheet = workbook["고객DB"]
        headers = [cell.value for cell in worksheet[1]]
        assert headers[:len(original_headers)] == original_headers
        assert set(FIELDS) <= set(headers)
        assert worksheet.max_row == expected_rows + 1
        assert worksheet.cell(2, headers.index("사업자등록번호") + 1).value == expected_business_no
        assert worksheet.cell(2, headers.index("사용자확장메모") + 1).value == "합성 보존 메모"
        assert worksheet.cell(2, headers.index("사용자계산식") + 1).value == "=1+2"
        assert worksheet.cell(2, headers.index("매출액") + 1).value == "=10+20"
        comment = worksheet.cell(2, headers.index("사용자확장메모") + 1).comment
        assert comment.text == "합성 보존 주석" and comment.author == "합성 작성자"
        for sheet_name in ("상시정책자금DB", "고용지원금DB", "코드표", "사용가이드", "사용자추가시트"):
            assert sheet_name in workbook.sheetnames
            assert workbook[sheet_name]["A1"].value == "합성 보존값 " + sheet_name
            assert workbook[sheet_name]["B2"].value == "=1+2"
    finally:
        workbook.close()


def test_manual_mobile_update_preserves_existing_workbook_data(legacy_workbook_with_extensions):
    path = legacy_workbook_with_extensions
    ok, _ = utils.update_user_customer_record("synthetic-owner", 0, {"대표자 휴대전화": MOBILE})
    assert ok
    _assert_existing_workbook_data_preserved(path)
    assert utils._read_cumulative_customer_db(path).iloc[0]["대표자 휴대전화"] == MOBILE


def test_new_customer_append_preserves_existing_workbook_data(legacy_workbook_with_extensions):
    path = legacy_workbook_with_extensions
    incoming = company(**{"업체명": "두번째 합성기업", "사업자등록번호": "-".join(("222", "22", "22222"))})
    _, count, _, _, _ = utils.append_cretop_to_user_customer_db("unused.pdf", "synthetic-owner", extracted_data=incoming)
    assert count == 1
    _assert_existing_workbook_data_preserved(path, expected_rows=2)
    assert utils._read_cumulative_customer_db(path).iloc[1]["대표자 휴대전화"] == MOBILE


def test_business_number_edit_is_allowed_without_losing_existing_workbook_data(legacy_workbook_with_extensions):
    changed_business_no = "-".join(("222", "22", "22222"))
    ok, _ = utils.update_user_customer_record(
        "synthetic-owner", 0, {"사업자등록번호": changed_business_no, "대표자 휴대전화": MOBILE},
    )
    assert ok
    _assert_existing_workbook_data_preserved(legacy_workbook_with_extensions, expected_business_no=changed_business_no)


@pytest.mark.parametrize("invalid_frame", ["empty", "fewer_rows", "reordered", "duplicate_index"])
def test_unsafe_customer_dataframe_never_rewrites_original_workbook(legacy_workbook_with_extensions, invalid_frame):
    from openpyxl import load_workbook

    path = legacy_workbook_with_extensions
    workbook = load_workbook(path)
    worksheet = workbook["고객DB"]
    headers = [cell.value for cell in worksheet[1]]
    other = company(**{"업체명": "두번째 합성기업", "사업자등록번호": "-".join(("222", "22", "22222"))})
    worksheet.append([other.get(column, "") for column in headers])
    workbook.save(path)
    workbook.close()
    frame = utils._read_cumulative_customer_db(path)
    assert len(frame) == 2
    if invalid_frame == "empty":
        incoming = frame.iloc[:0].copy()
    elif invalid_frame == "fewer_rows":
        incoming = frame.iloc[:1].copy()
    elif invalid_frame == "reordered":
        incoming = frame.iloc[::-1].reset_index(drop=True)
    else:
        incoming = frame.copy()
        incoming.index = [0, 0]
    before = path.read_bytes()
    previous_mtime = path.stat().st_mtime_ns
    with pytest.raises(ValueError):
        utils._write_cumulative_customer_db(path, incoming)
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == previous_mtime


def test_customer_read_failure_never_rewrites_original_workbook(legacy_workbook_with_extensions, monkeypatch):
    path = legacy_workbook_with_extensions
    before = path.read_bytes()
    previous_mtime = path.stat().st_mtime_ns
    monkeypatch.setattr(utils.pd, "read_excel", Mock(side_effect=ValueError("synthetic read failure")))
    ok, _ = utils.update_user_customer_record("synthetic-owner", 0, {"대표자 휴대전화": MOBILE})
    assert not ok
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == previous_mtime


def test_legacy_named_sheet_is_copied_without_losing_original_and_second_read_is_unchanged(legacy_workbook_with_extensions):
    from openpyxl import load_workbook

    path = legacy_workbook_with_extensions
    legacy_name = utils.LEGACY_CUMULATIVE_SHEET_NAME
    workbook = load_workbook(path, data_only=False)
    original = workbook["고객DB"]
    original.title = legacy_name
    original_values = list(original.values)
    original_headers = [cell.value for cell in original[1]]
    original_sheet_names = workbook.sheetnames[:]
    workbook.save(path)
    workbook.close()

    resolved, count, converted = utils.ensure_user_cumulative_db_format("synthetic-owner")
    assert resolved == path and count == 1 and converted
    _assert_existing_workbook_data_preserved(path)
    workbook = load_workbook(path, data_only=False)
    try:
        assert set(workbook.sheetnames) == set(original_sheet_names) | {"고객DB"}
        legacy = workbook[legacy_name]
        assert list(legacy.values) == original_values
        comment = legacy.cell(2, original_headers.index("사용자확장메모") + 1).comment
        assert comment.text == "합성 보존 주석" and comment.author == "합성 작성자"
    finally:
        workbook.close()

    before_second_read = path.read_bytes()
    previous_mtime = path.stat().st_mtime_ns
    assert utils.ensure_user_cumulative_db_format("synthetic-owner") == (path, 1, False)
    assert path.read_bytes() == before_second_read
    assert path.stat().st_mtime_ns == previous_mtime
